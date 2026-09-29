import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from app.schemas.paper_schema import PaperLayout, PaperSchema, Question
from app.services.pdf_service import marks_label, question_label

_PAGE_MM = {"A4": (210, 297), "Letter": (215.9, 279.4), "Legal": (215.9, 355.6)}
_MARGIN_IN = {"narrow": 0.5, "normal": 0.75, "wide": 1.0}
_FONT = {"serif": "Times New Roman", "sans": "Arial"}


def _accent(layout: PaperLayout):
    if layout.accent_color and re.fullmatch(r"#?[0-9a-fA-F]{6}", layout.accent_color.strip()):
        return RGBColor.from_string(layout.accent_color.strip().lstrip("#").upper())
    return None


def _border(element, tag: str, **attrs):
    """Attach a <w:pBdr>/<w:tblBorders>-style child with the given sides."""
    container = OxmlElement(tag)
    for side, spec in attrs.items():
        el = OxmlElement(f"w:{side}")
        for k, v in spec.items():
            el.set(qn(f"w:{k}"), v)
        container.append(el)
    element.append(container)


class DocxGenerationService:
    @classmethod
    def generate_question_paper_docx(cls, schema: PaperSchema, include_solutions: bool = False) -> bytes:
        """Render the paper as an editable Word document following schema.layout."""
        layout = schema.layout
        doc = cls._new_document(layout)
        cls._header(doc, schema)

        if layout.columns > 1:
            # Header stays full width; a continuous section break starts the question columns.
            section = doc.add_section(WD_SECTION.CONTINUOUS)
            cols = section._sectPr.xpath("./w:cols")
            cols_el = cols[0] if cols else OxmlElement("w:cols")
            cols_el.set(qn("w:num"), str(layout.columns))
            cols_el.set(qn("w:space"), str(int(Inches(0.3).twips)))
            cols_el.set(qn("w:sep"), "1")
            if not cols:
                section._sectPr.append(cols_el)

        first = doc.sections[0]
        usable = first.page_width - first.left_margin - first.right_margin
        col_width = int((usable - Inches(0.3) * (layout.columns - 1)) / layout.columns)
        for sec in schema.sections:
            cls._section_heading(doc, layout, sec.title, sec.instructions)
            for q in sec.questions:
                cls._question(doc, layout, q, include_solutions, col_width)
        return cls._save(doc)

    @classmethod
    def generate_solutions_docx(cls, schema: PaperSchema) -> bytes:
        """Answer key and step-by-step solutions only."""
        layout = schema.layout
        doc = cls._new_document(layout)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Answer Key & Solutions: {schema.metadata.title}")
        r.bold = True
        r.font.size = Pt(layout.font_size + 3)
        if _accent(layout):
            r.font.color.rgb = _accent(layout)
        for sec in schema.sections:
            cls._section_heading(doc, layout, sec.title, None)
            for q in sec.questions:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.add_run(f"{question_label(layout, q.question_number)} ").bold = True
                p.add_run("Answer: ").bold = True
                p.add_run(q.answer_key or "(no answer provided)")
                if q.detailed_solution:
                    sp = doc.add_paragraph()
                    sp.paragraph_format.left_indent = Inches(0.3)
                    sr = sp.add_run(q.detailed_solution)
                    sr.italic = True
                    sr.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return cls._save(doc)

    # ---------- building blocks ----------
    @staticmethod
    def _new_document(layout: PaperLayout) -> Document:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = _FONT[layout.font_family]
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT[layout.font_family])
        normal.font.size = Pt(layout.font_size)
        normal.paragraph_format.line_spacing = layout.line_spacing
        normal.paragraph_format.space_after = Pt(2)

        section = doc.sections[0]
        w, h = _PAGE_MM.get(layout.page_size, _PAGE_MM["A4"])
        if layout.orientation == "landscape":
            w, h = h, w
            section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(w), Mm(h)
        m = Inches(_MARGIN_IN.get(layout.margins, 0.75))
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = m

        footer = section.footer.paragraphs[0]
        if layout.footer_text:
            footer.add_run(layout.footer_text).font.size = Pt(max(layout.font_size - 2, 7))
        if layout.show_page_numbers:
            footer = section.footer.add_paragraph() if layout.footer_text else footer
            footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = footer.add_run("Page ")
            run.font.size = Pt(max(layout.font_size - 2, 7))
            for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
                fld_run = footer.add_run()
                if kind:
                    fld = OxmlElement("w:fldChar")
                    fld.set(qn("w:fldCharType"), kind)
                else:
                    fld = OxmlElement("w:instrText")
                    fld.set(qn("xml:space"), "preserve")
                    fld.text = text
                fld_run._r.append(fld)
        return doc

    @staticmethod
    def _header(doc: Document, schema: PaperSchema) -> None:
        meta, layout = schema.metadata, schema.layout
        align = WD_ALIGN_PARAGRAPH.CENTER if layout.header_alignment == "center" else WD_ALIGN_PARAGRAPH.LEFT
        color = _accent(layout)

        def heading(text, size, bold=True, italic=False):
            p = doc.add_paragraph()
            p.alignment = align
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            r.font.size = Pt(size)
            if color and bold:
                r.font.color.rgb = color

        if meta.institution_name:
            heading(meta.institution_name.upper(), layout.font_size + 4)
        heading(meta.title, layout.font_size + 3)
        if meta.subtitle:
            heading(meta.subtitle, layout.font_size, bold=False, italic=True)

        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        subject = meta.subject + (f" ({meta.class_grade})" if meta.class_grade else "")
        cells = [
            (0, 0, "Subject: ", subject, WD_ALIGN_PARAGRAPH.LEFT),
            (0, 1, "Maximum Marks: ", f"{meta.total_marks:g}", WD_ALIGN_PARAGRAPH.RIGHT),
            (1, 0, "Date: ", meta.date or "____________", WD_ALIGN_PARAGRAPH.LEFT),
            (1, 1, "Time Allowed: ", f"{meta.duration_minutes} Minutes", WD_ALIGN_PARAGRAPH.RIGHT),
        ]
        for row, col, label, value, cell_align in cells:
            p = table.cell(row, col).paragraphs[0]
            p.alignment = cell_align
            p.add_run(label).bold = True
            p.add_run(value)
        if layout.boxed_header:
            hexcolor = (layout.accent_color or "000000").lstrip("#") if color else "000000"
            spec = {"val": "single", "sz": "8", "space": "0", "color": hexcolor}
            _border(table._tbl.tblPr, "w:tblBorders", top=spec, left=spec, bottom=spec, right=spec)

        if layout.student_fields:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            for i, field in enumerate(layout.student_fields[:6]):
                if i:
                    p.add_run("      ")
                p.add_run(f"{field}: ").bold = True
                p.add_run("________________")

        if meta.general_instructions:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run("General Instructions:").bold = True
            for inst in meta.general_instructions:
                ip = doc.add_paragraph(style="List Bullet")
                ip.add_run(inst).font.size = Pt(layout.font_size - 0.5)

        divider = doc.add_paragraph()
        _border(divider._p.get_or_add_pPr(), "w:pBdr",
                bottom={"val": "single", "sz": "8", "space": "1", "color": "000000"})

    @staticmethod
    def _section_heading(doc: Document, layout: PaperLayout, title: str, instructions) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.bold = True
        r.font.size = Pt(layout.font_size + 1)
        if _accent(layout):
            r.font.color.rgb = _accent(layout)
        if instructions:
            ip = doc.add_paragraph()
            ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ip.paragraph_format.keep_with_next = True
            ir = ip.add_run(f"({instructions})")
            ir.italic = True
            ir.font.size = Pt(layout.font_size - 1)

    @staticmethod
    def _question(doc: Document, layout: PaperLayout, q: Question, include_solutions: bool, col_width: int) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = bool(q.options)
        p.add_run(f"{question_label(layout, q.question_number)} ").bold = True
        p.add_run(q.text)
        if layout.show_marks and q.marks:
            p.add_run(f"  {marks_label(q.marks)}").bold = True

        if q.options:
            opts = [(f"({o.label}) ", o.text) for o in q.options]
            if layout.option_layout == "inline":
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Inches(0.3)
                for i, (label, text) in enumerate(opts):
                    if i:
                        op.add_run("     ")
                    op.add_run(label).bold = True
                    op.add_run(text)
            elif layout.option_layout == "grid":
                rows = (len(opts) + 1) // 2
                grid = doc.add_table(rows=rows, cols=2)
                grid.autofit = False
                for column in grid.columns:
                    column.width = col_width // 2
                for i, (label, text) in enumerate(opts):
                    cell = grid.cell(i // 2, i % 2)
                    cell.width = col_width // 2
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.left_indent = Inches(0.2)
                    cp.add_run(label).bold = True
                    cp.add_run(text)
            else:
                for label, text in opts:
                    op = doc.add_paragraph()
                    op.paragraph_format.left_indent = Inches(0.3)
                    op.paragraph_format.space_after = Pt(0)
                    op.add_run(label).bold = True
                    op.add_run(text)
        elif layout.answer_lines and not include_solutions:
            for _ in range(layout.answer_lines):
                line = doc.add_paragraph()
                line.paragraph_format.space_before = Pt(layout.font_size)
                _border(line._p.get_or_add_pPr(), "w:pBdr",
                        bottom={"val": "single", "sz": "4", "space": "1", "color": "999999"})

        if include_solutions and (q.answer_key or q.detailed_solution):
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Inches(0.3)
            if q.answer_key:
                r = sp.add_run(f"Answer: {q.answer_key}")
                r.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            if q.detailed_solution:
                r = sp.add_run(("\n" if q.answer_key else "") + f"Solution: {q.detailed_solution}")
                r.italic = True
                r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    @staticmethod
    def _save(doc: Document) -> bytes:
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
