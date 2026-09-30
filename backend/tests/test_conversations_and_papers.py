import io
import json
import time

from docx import Document


def wait_job(client, headers, job_id, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        job = client.get(f"/api/v1/jobs/{job_id}", headers=headers).json()
        if job["status"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def chat(client, headers, conv_id, text):
    """Send a message, wait for the assistant, return (job, last assistant message, its metadata)."""
    res = client.post(f"/api/v1/conversations/{conv_id}/messages", json={"content": text}, headers=headers)
    assert res.status_code == 200, res.text
    meta = json.loads(res.json()["metadata_json"])
    assert meta["action"] == "working"
    job = wait_job(client, headers, meta["job_id"])
    msgs = client.get(f"/api/v1/conversations/{conv_id}", headers=headers).json()["messages"]
    last = msgs[-1]
    return job, last, json.loads(last["metadata_json"] or "{}")


def generate(client, headers, conv_id, text="generate"):
    job, _, meta = chat(client, headers, conv_id, text)
    assert job["status"] == "done", job
    assert meta["action"] == "paper_ready"
    return meta["paper_id"]


def new_chat(client, headers, title="Test"):
    return client.post("/api/v1/conversations", json={"title": title}, headers=headers).json()["id"]


def test_create_and_list_conversations(client, auth_headers):
    res = client.post("/api/v1/conversations", json={"title": "NEET Physics Test"}, headers=auth_headers)
    assert res.status_code == 201
    conv_id = res.json()["id"]
    assert any(c["id"] == conv_id for c in client.get("/api/v1/conversations", headers=auth_headers).json())


def test_full_generate_edit_export_flow(client, auth_headers):
    conv_id = new_chat(client, auth_headers, "Class 12 Physics Test")

    _, reply, _ = chat(client, auth_headers, conv_id, "I want to create a physics paper on electrostatics.")
    assert "difficulty" in reply["content"].lower()

    paper_id = generate(client, auth_headers, conv_id, "Please generate the paper now with 35 marks, moderate difficulty.")
    detail = client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).json()
    assert detail["latest_paper_id"] == paper_id and detail["active_job_id"] is None

    paper = client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()
    assert paper["kind"] == "paperly" and paper["metadata"]["subject"] == "Physics"
    assert paper["conversation_id"] == conv_id

    edit_res = client.post(f"/api/v1/papers/{paper_id}/edit", json={"instruction": "Make question 1 harder."}, headers=auth_headers)
    assert edit_res.status_code == 202
    edit_job = wait_job(client, auth_headers, edit_res.json()["job_id"])
    assert edit_job["status"] == "done"
    assert edit_job["result"]["version_number"] == 2
    assert "Question 1" in edit_job["result"]["change_summary"]

    for path in ("export/docx", "export/pdf", "solutions/docx", "solutions/pdf"):
        res = client.get(f"/api/v1/papers/{paper_id}/{path}", headers=auth_headers)
        assert res.status_code == 200 and len(res.content) > 500, path

    versions = client.get(f"/api/v1/papers/{paper_id}/versions", headers=auth_headers).json()
    assert [v["version_number"] for v in versions][:2] == [2, 1]
    assert client.post(f"/api/v1/papers/{paper_id}/revert/1", headers=auth_headers).json()["version_number"] == 1


def test_regeneration_creates_new_version_and_lists_papers(client, auth_headers):
    conv_id = new_chat(client, auth_headers, "Regen")
    paper_id = generate(client, auth_headers, conv_id)
    assert generate(client, auth_headers, conv_id, "generate again") == paper_id
    versions = client.get(f"/api/v1/papers/{paper_id}/versions", headers=auth_headers).json()
    assert [v["version_number"] for v in versions] == [2, 1]
    papers = client.get("/api/v1/papers", headers=auth_headers).json()
    assert any(p["id"] == paper_id and p["question_count"] == 3 and p["kind"] == "paperly" for p in papers)


def test_question_about_paper_gets_answer_without_changing_it(client, auth_headers):
    conv_id = new_chat(client, auth_headers)
    paper_id = generate(client, auth_headers, conv_id)
    _, reply, meta = chat(client, auth_headers, conv_id, "Why is the max marks showing this number?")
    assert meta["action"] == "reply"
    versions = client.get(f"/api/v1/papers/{paper_id}/versions", headers=auth_headers).json()
    assert len(versions) == 1  # nothing was changed


def test_chat_edits_paperly_layout(client, auth_headers):
    conv_id = new_chat(client, auth_headers)
    paper_id = generate(client, auth_headers, conv_id)
    job, _, meta = chat(client, auth_headers, conv_id, "Make it two columns with a serif font")
    assert job["status"] == "done" and meta["action"] == "paper_updated" and meta["version"] == 2
    paper = client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()
    assert paper["layout"]["columns"] == 2 and paper["layout"]["font_family"] == "serif"


def _sample_docx() -> bytes:
    doc = Document()
    title = doc.add_paragraph()
    title.add_run("Class 8 Science - Unit Test").bold = True
    doc.add_paragraph("Maximum Marks: 70")
    for n, q in enumerate(["What is photosynthesis?", "Name the parts of a flower.", "What is friction?"], 1):
        p = doc.add_paragraph()
        p.add_run(f"Q{n}.").bold = True
        p.add_run(f" {q}")
        p.add_run(" [2]").bold = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_uploaded_word_document_is_edited_in_place(client, auth_headers):
    conv_id = new_chat(client, auth_headers, "Edit doc")
    up = client.post(f"/api/v1/conversations/{conv_id}/upload",
                     files={"file": ("science_test.docx", _sample_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                     headers=auth_headers)
    assert up.status_code == 200
    detail = client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).json()
    paper_id = detail["latest_paper_id"]
    assert paper_id, "uploading a Word file should make it the chat's document"
    assert json.loads(detail["messages"][-1]["metadata_json"])["action"] == "paper_ready"

    job, reply, meta = chat(client, auth_headers, conv_id, "Change the last question to What is inertia?")
    assert job["status"] == "done", job
    assert meta["action"] == "paper_updated" and meta["changes"][0]["after"] == "Q3. What is inertia? [2]"

    paper = client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()
    assert paper["kind"] == "docx"
    texts = [b["text"] for b in paper["blocks"] if b["text"].strip()]
    assert texts == ["Class 8 Science - Unit Test", "Maximum Marks: 70", "Q1. What is photosynthesis? [2]",
                     "Q2. Name the parts of a flower. [2]", "Q3. What is inertia? [2]"]

    exported = client.get(f"/api/v1/papers/{paper_id}/export/docx", headers=auth_headers)
    assert exported.status_code == 200
    runs = Document(io.BytesIO(exported.content)).paragraphs[-1].runs
    assert runs[0].text == "Q3." and runs[0].bold  # bold question label kept
    assert runs[-1].text == " [2]" and runs[-1].bold  # bold marks kept

    # Word stays Word; no silent re-creation.
    assert client.get(f"/api/v1/papers/{paper_id}/export/pdf", headers=auth_headers).status_code in (200, 400)
    assert client.get(f"/api/v1/papers/{paper_id}/solutions/pdf", headers=auth_headers).status_code == 400


def test_uploaded_pdf_is_edited_in_place(client, auth_headers):
    from app.schemas.paper_schema import PaperSchema, PaperMetadata, Section, Question
    from app.services.pdf_service import PdfGenerationService

    schema = PaperSchema(
        metadata=PaperMetadata(title="Unit Test", subject="Science", total_marks=70),
        sections=[Section(id="s", title="Questions", questions=[
            Question(id=f"q{i}", question_number=i, type="short_answer", text=t, marks=2)
            for i, t in enumerate(["What is photosynthesis?", "What is friction?"], 1)
        ])],
    )
    pdf = PdfGenerationService.generate_pdf_bytes(schema)
    conv_id = new_chat(client, auth_headers, "PDF")
    assert client.post(f"/api/v1/conversations/{conv_id}/upload",
                       files={"file": ("unit_test.pdf", pdf, "application/pdf")}, headers=auth_headers).status_code == 200
    paper_id = client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).json()["latest_paper_id"]

    job, _, meta = chat(client, auth_headers, conv_id, "Change the last question to What is inertia?")
    assert job["status"] == "done", job
    assert meta["action"] == "paper_updated"

    paper = client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()
    assert paper["kind"] == "pdf" and paper["page_count"] == 1
    page = client.get(f"/api/v1/papers/{paper_id}/pages/1", headers=auth_headers)
    assert page.status_code == 200 and page.content[:4] == b"\x89PNG"

    from app.services.doc_edit import pdf_blocks
    edited = client.get(f"/api/v1/papers/{paper_id}/export/pdf", headers=auth_headers).content
    texts = [b.text for b in pdf_blocks(edited)]
    assert any("What is inertia?" in t for t in texts)
    assert any("What is photosynthesis?" in t for t in texts)
    assert not any("friction" in t for t in texts)
    assert any("Max Marks: 70" in t for t in texts)


def test_max_marks_are_never_overwritten():
    from app.schemas.paper_schema import PaperSchema, PaperMetadata, Section, Question
    from app.services.ai_service import AIService

    # e.g. internal "OR" choices: question marks add up to more than the paper's maximum.
    schema = PaperSchema(
        metadata=PaperMetadata(title="T", total_marks=70),
        sections=[Section(id="s", title="A", questions=[
            Question(id=f"q{i}", question_number=i, text="x", marks=10) for i in range(1, 12)
        ])],
    )
    assert AIService._normalize(schema).metadata.total_marks == 70
    edited, _ = AIService._mock_paper_edit(schema, "Make Q1 harder")
    assert edited.metadata.total_marks == 70


def test_delete_conversation(client, auth_headers):
    conv_id = new_chat(client, auth_headers, "Temp")
    assert client.delete(f"/api/v1/conversations/{conv_id}", headers=auth_headers).status_code == 204
    assert client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).status_code == 404


def test_upload_reference_and_reject_unsupported(client, auth_headers):
    conv_id = new_chat(client, auth_headers)
    ok = client.post(f"/api/v1/conversations/{conv_id}/upload",
                     files={"file": ("notes.txt", b"Chapter 1: Newton's laws", "text/plain")}, headers=auth_headers)
    assert ok.status_code == 200
    assert client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).json()["latest_paper_id"] is None
    bad = client.post(f"/api/v1/conversations/{conv_id}/upload",
                      files={"file": ("archive.zip", b"PK\x03\x04", "application/zip")}, headers=auth_headers)
    assert bad.status_code == 415
    photo = client.post(f"/api/v1/conversations/{conv_id}/upload",
                        files={"file": ("photo.jpg", b"\xff\xd8\xff", "image/jpeg")}, headers=auth_headers)
    assert photo.status_code == 422  # photos need Gemini, which tests don't have


def test_ai_failure_posts_error_message(client, auth_headers, monkeypatch):
    from app.services.ai_service import AIService, AIServiceError

    async def boom(*args, **kwargs):
        raise AIServiceError("Gemini rejected the API key.")

    monkeypatch.setattr(AIService, "chat_turn", boom)
    conv_id = new_chat(client, auth_headers)
    job, last, meta = chat(client, auth_headers, conv_id, "hi")
    assert job["status"] == "error" and "API key" in job["error"]
    assert meta["action"] == "error" and "API key" in last["content"]


def test_second_upload_is_readable_reference_and_can_be_opened(client, auth_headers, monkeypatch):
    """Word paper first, then a PDF: the Word paper stays open, the assistant receives the PDF to read
    (instead of asking the teacher to paste it), and 'edit the PDF instead' switches documents."""
    from app.schemas.paper_schema import PaperSchema, PaperMetadata, Section, Question
    from app.services.ai_service import AIService, AgentTurn
    from app.services.pdf_service import PdfGenerationService

    conv_id = new_chat(client, auth_headers, "Two files")
    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    client.post(f"/api/v1/conversations/{conv_id}/upload",
                files={"file": ("my_paper.docx", _sample_docx(), docx_mime)}, headers=auth_headers)
    pdf = PdfGenerationService.generate_pdf_bytes(PaperSchema(
        metadata=PaperMetadata(title="Other", total_marks=5),
        sections=[Section(id="s", title="Chapter 14", questions=[Question(id="q1", question_number=1, text="What is sound?")])],
    ))
    client.post(f"/api/v1/conversations/{conv_id}/upload",
                files={"file": ("chapter14.pdf", pdf, "application/pdf")}, headers=auth_headers)

    detail = client.get(f"/api/v1/conversations/{conv_id}", headers=auth_headers).json()
    paper_id = detail["latest_paper_id"]
    assert client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()["kind"] == "docx"
    assert json.loads(detail["messages"][-1]["metadata_json"])["action"] == "file_received"

    seen = {}

    async def fake_turn(cls, history, doc_kind, doc_text, sources):
        seen["doc_kind"], seen["sources"] = doc_kind, [(s.filename, bool(s.raw_path)) for s in sources]
        if "instead" in history[-1]["content"]:
            return AgentTurn(reply="Now editing chapter14.pdf.", action="open_file", source_file="chapter14.pdf")
        return AgentTurn(reply="ok")

    monkeypatch.setattr(AIService, "chat_turn", classmethod(fake_turn))
    chat(client, auth_headers, conv_id, "Replace chapter 10 with the chapter 14 questions from the PDF")
    assert seen["doc_kind"] == "docx"
    assert seen["sources"] == [("chapter14.pdf", True)]  # the PDF is given to the assistant; the open Word file isn't

    job, _, meta = chat(client, auth_headers, conv_id, "Edit chapter14.pdf instead")
    assert job["status"] == "done" and meta["opened_file"] == "chapter14.pdf"
    assert client.get(f"/api/v1/papers/{paper_id}", headers=auth_headers).json()["kind"] == "pdf"
    chat(client, auth_headers, conv_id, "anything")
    assert seen["sources"] == [("my_paper.docx", True)]  # now the Word file is the "other" file
